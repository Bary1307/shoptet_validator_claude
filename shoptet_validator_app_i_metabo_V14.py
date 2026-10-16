import streamlit as st
from lxml import etree
import pandas as pd
import json
import re
from copy import deepcopy
from urllib.parse import quote

st.set_page_config(page_title="Univerzální Feed → Shoptet Validator", layout="wide")

# =============================================================================
# SLOVNÍK SHOPTET TAGŮ
# =============================================================================
SHOPTET_DICT = {
    "CODE": {"desc": "Unikátní kód / SKU produktu.", "mandatory": True},
    "NAME": {"desc": "Název produktu.", "mandatory": True},
    "SHORT_DESCRIPTION": {"desc": "Krátký popis.", "mandatory": False},
    "DESCRIPTION": {"desc": "Detailní popis produktu, podporuje HTML.", "mandatory": False},
    "MANUFACTURER": {"desc": "Značka / výrobce.", "mandatory": False},
    "EAN": {"desc": "EAN / čárový kód.", "mandatory": False},
    "PART_NUMBER": {"desc": "Katalogové číslo výrobce.", "mandatory": False},
    "CURRENCY": {"desc": "Měna ceny (CZK, EUR...).", "mandatory": True},
    "PRICE": {"desc": "Cena produktu.", "mandatory": True},
    "VAT": {"desc": "Sazba DPH v procentech (např. 21).", "mandatory": False},
    "STANDARD_PRICE": {"desc": "Původní/přeškrtnutá cena před naší slevou (zobrazí se v Shoptetu jako přeškrtnutá).", "mandatory": False},
    "CATEGORY": {"desc": "Cesta kategorie.", "mandatory": True},
    "IMAGE": {"desc": "Přímá veřejná URL adresa obrázku.", "mandatory": False},
    "WEIGHT": {"desc": "Hmotnost v kg.", "mandatory": False},
    "WIDTH": {"desc": "Šířka produktu v cm.", "mandatory": False},
    "HEIGHT": {"desc": "Výška produktu v cm.", "mandatory": False},
    "DEPTH": {"desc": "Hloubka/délka produktu v cm (element LENGTH Shoptet nepovoluje).", "mandatory": False},
    "VISIBILITY": {"desc": "'visible' nebo 'hidden'.", "mandatory": False},
    "AVAILABILITY_IN_STOCK": {"desc": "Text dostupnosti, když je produkt skladem (např. 'Skladem').", "mandatory": False},
    "AVAILABILITY_OUT_OF_STOCK": {"desc": "Text dostupnosti, když produkt není skladem - musí odpovídat jedné z předdefinovaných hodnot Shoptetu.", "mandatory": False},
    "PARAMETERS/PARAMETER": {"desc": "POUZE pro velikosti (oblečení/obuv).", "mandatory": False},
    "INFORMATION_PARAMETERS/INFORMATION_PARAMETER": {"desc": "Obecné/technické parametry produktu.", "mandatory": False},
}

DEFAULT_PROFILE = {
    "profile_name": "Nová šablona",
    "item_tag": "",
    "fields": {t: {"source": ""} for t in ["CODE", "NAME", "DESCRIPTION", "MANUFACTURER", "EAN", "PART_NUMBER"]},
    "price": {"mode": "direct", "source": "", "price_type_attr": "", "price_type_value": ""},
    "currency": "CZK",
    "category": {"mode": "fixed", "fixed_value": "Nezařazeno", "source": "", "sep": " > "},
    "weight": {"mode": "none", "source": "", "feature_element": "", "fname_field": "", "fvalue_field": "",
               "funit_field": "", "names": "Hmotnost, Weight",
               "width_names": "", "height_names": "", "length_names": ""},
    "images": {"mode": "none", "source": "", "mime_element": "", "purpose_field": "", "purpose_value": "normal",
               "source_field": "", "max_images": 8},
    "params": {"repeat_element": "", "classifier_scope": "", "classifier_field": "", "fname_field": "", "fvalue_field": "",
               "funit_field": "", "value_map": {}, "marketing_target": "SHORT_DESCRIPTION"},
    "visibility_default": "visible",
    "videos": {"mode": "none", "mime_element": "MIME", "type_field": "MIME_TYPE", "type_value": "url",
               "descr_field": "MIME_DESCR", "descr_match": "video", "source_field": "MIME_SOURCE",
               "wrapper_tag": "RELATED_VIDEOS", "add_text": False},
    "references": {"enabled": False, "reference_element": "ARTICLE_REFERENCE", "type_attr": "type",
                   "id_field": "ART_ID_TO", "included_types": [], "wrapper_tag": "RELATED_PRODUCTS", "item_tag": "CODE"},
    "name_prefix": "",
    "name_suffix": "",
    "description_format": {"enabled": False, "paragraphs": True, "extract_packaging": True,
                            "headings": True, "bold_numbers": True, "dedupe": True},
    "name_cleanup": {"enabled": False, "strip_after_semicolon": True, "strip_own_code_parens": True},
    "text_tools": {"enabled": False, "collapse_spaces": True, "name_case": "none", "short_desc_max": 0,
                   "replace_rules": []},
    "price_adjust": {"enabled": False, "source_mode": "net", "discount_percent": 0, "vat_percent": "21"},
    "price_precision": "2dp",
    "availability": {"enabled": False, "in_stock": "Skladem", "out_of_stock": "2-3 dny"},
}

STEPS = [
    "1. Nahrání feedu",
    "2. Výběr produktů",
    "3. Základní mapování polí",
    "4. Kategorie",
    "5. Parametry",
    "6. Obrázky",
    "7. Kontrola a ruční opravy",
    "8. Export a šablona",
]


# =============================================================================
# POMOCNÉ FUNKCE
# =============================================================================
def gtext(el, tag):
    if not tag:
        return None
    r = el.xpath(f"./*[local-name()='{tag}']")
    return r[0].text.strip() if r and r[0].text else None


def gtext_deep(el, tag):
    if not tag:
        return None
    r = el.xpath(f".//*[local-name()='{tag}']")
    return r[0].text.strip() if r and r[0].text else None


def gall_deep(el, tag):
    if not tag:
        return []
    return el.xpath(f".//*[local-name()='{tag}']")


def parse_xml_bytes(data):
    parser = etree.XMLParser(recover=True, huge_tree=True)
    return etree.fromstring(data, parser=parser)


def detect_feed_type(root):
    return "shoptet" if root.tag.split("}")[-1] == "SHOP" else "other"


def local_tag(el):
    return el.tag.split("}")[-1] if isinstance(el.tag, str) else str(el.tag)


def candidate_item_tags(root, min_count=2):
    counts, parent_refs, child_tag_sets = {}, {}, {}
    for el in root.iter():
        t = local_tag(el)
        counts[t] = counts.get(t, 0) + 1
        parent = el.getparent()
        if parent is not None:
            parent_refs.setdefault(t, set()).add(parent)
        s = child_tag_sets.setdefault(t, set())
        for child in el:
            s.add(local_tag(child))
    candidates = []
    for t, c in counts.items():
        if c < min_count or len(child_tag_sets.get(t, set())) < 2 or t == local_tag(root):
            continue
        n_parents = max(len(parent_refs.get(t, set())), 1)
        ratio = c / n_parents
        if ratio < 1.5:
            continue
        candidates.append((t, c, ratio, len(child_tag_sets.get(t, set()))))
    candidates.sort(key=lambda x: (-x[3], -x[2], -x[1]))
    return [(t, c) for t, c, _, _ in candidates]


def descendant_tags(item_el):
    tags = sorted(set(local_tag(e) for e in item_el.iter()))
    return [t for t in tags if t != local_tag(item_el)]


def aggregate_descendant_tags(items, sample_size=30):
    """Sjednotí podřazené tagy napříč VÍC produkty, ne jen prvním - různé produkty
    ve stejném feedu můžou mít různě neúplnou strukturu (chybějící cena, jiné MIME...)."""
    tags = set()
    for item in items[:sample_size]:
        tags.update(local_tag(e) for e in item.iter())
    tags.discard(local_tag(items[0])) if items else None
    return sorted(tags)


def repeat_children_candidates(item_el, min_count=2):
    """Najde opakující se podřazené elementy uvnitř JEDNOHO produktu (kandidáty na
    'jeden parametr/vlastnost'), např. FEATURE."""
    counts, child_tag_sets = {}, {}
    for el in item_el.iter():
        if el is item_el:
            continue
        t = local_tag(el)
        counts[t] = counts.get(t, 0) + 1
        s = child_tag_sets.setdefault(t, set())
        for child in el:
            s.add(local_tag(child))
    candidates = [(t, c) for t, c in counts.items() if c >= min_count and len(child_tag_sets.get(t, set())) >= 1]
    candidates.sort(key=lambda x: -x[1])
    return candidates


def classifier_field_candidates(root, item_tag, repeat_tag, sample_size=30):
    """Najde kandidáty na 'rozlišovací pole' druhu parametru - buď uvnitř samotné
    opakující se vlastnosti (repeat_tag, např. FDESCR uvnitř FEATURE), NEBO na jejím
    přímém rodiči (např. REFERENCE_FEATURE_GROUP_NAME na obalu ARTICLE_FEATURES,
    jak to dělá třeba Metabo). Vrací (scope, field) dvojice, scope je 'self'/'parent'."""
    items = root.xpath(f"//*[local-name()='{item_tag}']")[:sample_size]
    self_values, parent_values = {}, {}
    for item in items:
        for feat in item.xpath(f".//*[local-name()='{repeat_tag}']"):
            for child in feat:
                fname = local_tag(child)
                if child.text and child.text.strip():
                    self_values.setdefault(fname, []).append(child.text.strip())
            parent = feat.getparent()
            if parent is not None:
                for child in parent:
                    if local_tag(child) == repeat_tag:
                        continue
                    fname = local_tag(child)
                    if child.text and child.text.strip():
                        parent_values.setdefault(fname, []).append(child.text.strip())

    def score_all(values_dict):
        scored = []
        for f, vals in values_dict.items():
            distinct = len(set(vals))
            if 1 < distinct <= max(20, len(vals) * 0.3):
                scored.append((f, distinct))
        scored.sort(key=lambda x: x[1])
        return [f for f, _ in scored]

    # rodičovské (skupinové) pole preferujeme jako první tipy, protože je to častější
    # a spolehlivější vzor (jedna skupina = jeden druh parametru) než pole uvnitř samotné vlastnosti
    result = [("parent", f) for f in score_all(parent_values)]
    result += [("self", f) for f in score_all(self_values)]
    return result


YOUTUBE_ID_RE = re.compile(r"(?:youtube\.com/(?:watch\?(?:[^#]*&)?v=|embed/|shorts/|v/)|youtu\.be/)([A-Za-z0-9_-]{11})")


def extract_youtube_id(url):
    """Shoptet chce u videa jen kód YouTube videa (YOUTUBE_VIDEO_CODE), ne celou URL."""
    m = YOUTUBE_ID_RE.search(url or "")
    return m.group(1) if m else None


def get_attribute_values(root, item_tag, element_local_name, attr_name, sample_size=100):
    """Zjistí, jaké hodnoty atributu (např. type=...) se u daného elementu v praxi vyskytují -
    pro auto-detekci typů ARTICLE_REFERENCE (accessories/cross_selling apod.)."""
    items = root.xpath(f"//*[local-name()='{item_tag}']")[:sample_size]
    counts = {}
    for it in items:
        for el in it.xpath(f".//*[local-name()='{element_local_name}']"):
            v = el.get(attr_name)
            if v:
                counts[v] = counts.get(v, 0) + 1
    return sorted(counts.items(), key=lambda x: -x[1])


def get_classifier_values(root, item_tag, repeat_tag, scope, classifier_field, sample_size=200):
    items = root.xpath(f"//*[local-name()='{item_tag}']")[:sample_size]
    counts = {}
    for item in items:
        for feat in item.xpath(f".//*[local-name()='{repeat_tag}']"):
            if scope == "parent":
                parent = feat.getparent()
                v = gtext(parent, classifier_field) if parent is not None else None
            else:
                v = gtext(feat, classifier_field)
            if v:
                counts[v] = counts.get(v, 0) + 1
    return sorted(counts.items(), key=lambda x: -x[1])


def guess_bucket(value):
    v = (value or "").lower()
    if "market" in v:
        return "marketing"
    if "scope" in v or "deliver" in v or "lieferumfang" in v or "obsah bal" in v:
        return "packaging"
    if "techni" in v or "technic" in v:
        return "technical"
    if "velikost" in v or "size" in v:
        return "size"
    return "ignore"


# =============================================================================
# FORMÁTOVÁNÍ DESCRIPTION - z jednoho souvislého bloku textu udělá strukturu
# =============================================================================
PACKAGING_RE = re.compile(r"(balen[ií]|krabice|karton|obal(?:u|em)?|\d+\s*[×x]\s*\S)", re.IGNORECASE)
NUMBER_UNIT_RE = re.compile(
    r"(?<![A-Za-z0-9])(\d[\d\s\u00a0.,]*\d|\d)\s?(W|kg|g|mm|cm|km|V|Ah|mAh|°C|°F|l|ml|Nm|min|kWh)(?![A-Za-z0-9])"
)
SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?])\s+(?=[A-ZÁČĎÉĚÍŇÓŘŠŤÚŮÝŽ0-9])")


def clean_name_text(text, code_value, opts):
    """Očistí NAME od interního kódu v závorce (přesně shodného s CODE produktu) a od
    textu za středníkem (typicky typ balení), pokud si to appka nechá zapnuté. Např.
    'Aku utahovák SSW 18 LTX 550 BL (602404850); Lepenková krabice' + CODE=602404850
    -> 'Aku utahovák SSW 18 LTX 550 BL'."""
    if not text or not opts or not opts.get("enabled"):
        return text
    t = text
    if opts.get("strip_after_semicolon"):
        t = t.split(";")[0]
    if opts.get("strip_own_code_parens") and code_value:
        t = re.sub(r"\s*\(\s*" + re.escape(code_value) + r"\s*\)", "", t)
    return t.strip()


def apply_text_rules(text, field, cfg):
    """Obecné úpravy textu pole (NAME / SHORT_DESCRIPTION / DESCRIPTION) podle cfg['text_tools']."""
    if not text or not cfg or not cfg.get("enabled"):
        return text
    t = text
    for rule in cfg.get("replace_rules", []):
        f = (rule.get("field") or "").strip()
        find = rule.get("find") or ""
        if not find or f not in ("", "VŠECHNA", field):
            continue
        repl = rule.get("replace") or ""
        try:
            if rule.get("regex"):
                t = re.sub(find, repl, t)
            else:
                t = t.replace(find, repl)
        except re.error:
            pass
    if cfg.get("collapse_spaces"):
        t = re.sub(r"[ \t\u00a0]+", " ", t)
        t = re.sub(r" *\n *", "\n", t).strip()
    if field == "NAME":
        mode = cfg.get("name_case", "none")
        if mode == "first_upper" and t:
            t = t[0].upper() + t[1:]
        elif mode == "upper":
            t = t.upper()
        elif mode == "lower":
            t = t.lower()
        elif mode == "sentence" and t:
            t = t[0].upper() + t[1:].lower()
    if field == "SHORT_DESCRIPTION":
        mx = int(cfg.get("short_desc_max") or 0)
        if mx > 0:
            plain = re.sub(r"<[^>]+>", " ", t)
            plain = re.sub(r"\s+", " ", plain).strip()
            if len(plain) > mx:
                cut = plain[:mx].rsplit(" ", 1)[0].rstrip(",;:.- ")
                t = cut + "…"
            else:
                t = plain
    return t


def apply_text_tools(shop, profile):
    cfg = profile.get("text_tools")
    if not cfg or not cfg.get("enabled"):
        return
    for item in shop.iterchildren("SHOPITEM"):
        for field in ("NAME", "SHORT_DESCRIPTION", "DESCRIPTION"):
            el = item.find(field)
            if el is not None and el.text:
                el.text = apply_text_rules(el.text, field, cfg)


def _norm_words(t):
    t = re.sub(r"<[^>]+>", " ", t or "").lower()
    return re.findall(r"[0-9a-záčďéěíňóřšťúůýž]+", t)


def _stems(t):
    return {w[:5] for w in _norm_words(t) if len(w) > 2 or w.isdigit()}


def is_duplicate_text(candidate, reference_sentences, threshold=0.8):
    """True, pokud je candidate (odrážka) v podstatě už obsažen v některé z vět reference.
    Porovnává se podle kmenů slov (prvních 5 znaků), takže zvládne i české skloňování."""
    cw = " ".join(_norm_words(candidate))
    if not cw:
        return True
    cs = _stems(candidate)
    for ref in reference_sentences:
        rw = " ".join(_norm_words(ref))
        if cw in rw:
            return True
        if len(cs) >= 3:
            rs = _stems(ref)
            if rs and len(cs & rs) / len(cs) >= threshold:
                return True
    return False


def dedupe_list(items, reference_sentences, threshold=0.8):
    out, kept = [], []
    for it in items:
        if is_duplicate_text(it, reference_sentences, threshold) or is_duplicate_text(it, kept, 1.01):
            continue
        out.append(it)
        kept.append(it)
    return out


def format_description_html(text, opts, bullets=None, packaging_items=None):
    """Složí DESCRIPTION: popis (formátovaný), odrážky (marketing), a úplně na konci obsah balení.
    Pokud je zapnutá kontrola duplicit, odrážky/balení, které už jsou obsažené v textu, se vynechají."""
    opts = opts or {}
    fmt = bool(opts.get("enabled"))
    bullets = list(bullets or [])
    packaging_items = list(packaging_items or [])
    text = (text or "").strip()
    if not text and not bullets and not packaging_items:
        return text
    if not fmt and not bullets and not packaging_items:
        return text

    sentences = [x for x in SENTENCE_SPLIT_RE.split(text) if x.strip()] if text else []
    packaging_sent = []
    if fmt and opts.get("extract_packaging") and sentences:
        keep = []
        for c in sentences:
            (packaging_sent if PACKAGING_RE.search(c) else keep).append(c)
        sentences = keep

    if opts.get("dedupe", True):
        seen, uniq = [], []
        for c in sentences:
            if is_duplicate_text(c, seen, 1.01):
                continue
            uniq.append(c)
            seen.append(c)
        sentences = uniq
        bullets = dedupe_list(bullets, sentences)
        packaging_items = dedupe_list(packaging_items, packaging_sent + sentences)

    body = " ".join(sentences).strip()
    if fmt and opts.get("bold_numbers") and body:
        body = NUMBER_UNIT_RE.sub(r"<strong>\1 \2</strong>", body)

    if fmt and opts.get("paragraphs") and sentences:
        sents = [x for x in SENTENCE_SPLIT_RE.split(body) if x.strip()]
        paras = [" ".join(sents[i:i + 2]) for i in range(0, len(sents), 2)]
        body_html = "".join(f"<p>{p}</p>" for p in paras)
    else:
        body_html = f"<p>{body}</p>" if (fmt and body) else body

    html = ""
    if fmt and opts.get("headings") and body:
        html += "<h3>Popis</h3>"
    html += body_html
    if bullets:
        if fmt and opts.get("headings"):
            html += "<h3>Vlastnosti</h3>"
        html += "<ul>" + "".join(f"<li>{b}</li>" for b in bullets) + "</ul>"
    pk = []
    if packaging_sent:
        pk.append("<p>" + " ".join(x.strip() for x in packaging_sent) + "</p>")
    if packaging_items:
        pk.append("<ul>" + "".join(f"<li>{b}</li>" for b in packaging_items) + "</ul>")
    if pk:
        html += "<h3>Obsah balení</h3>" + "".join(pk)   # vždy úplně na konci
    return html


# =============================================================================
# KATEGORIE - stromový resolver (pro BMEcat styl, volitelný)
# =============================================================================
def build_tree_category_map(root, cfg):
    cats = {}
    for cs in root.xpath(f"//*[local-name()='{cfg['struct_element']}']"):
        gid = gtext(cs, cfg["group_id"])
        cats[gid] = {"name": gtext(cs, cfg["group_name"]), "parent": gtext(cs, cfg["parent_id"]), "type": cs.get("type")}
    return cats


def build_tree_links(root, cfg):
    m = {}
    for link in root.xpath(f"//*[local-name()='{cfg['link_element']}']"):
        art_id = gtext(link, cfg["id_field"])
        gid = gtext(link, cfg["group_id_field"])
        if art_id:
            m.setdefault(art_id, []).append(gid)
    return m


def tree_category_path(gid, cats, sep, root_type_value):
    chain, seen, cur = [], set(), gid
    while cur and cur in cats and cur not in seen:
        seen.add(cur)
        node = cats[cur]
        if node["type"] != root_type_value:
            chain.append(node["name"])
        cur = node["parent"]
    chain.reverse()
    return sep.join([c for c in chain if c])


# =============================================================================
# APLIKACE PROFILU -> SHOP/SHOPITEM (čistá funkce, žádný vedlejší stav)
# =============================================================================
def apply_profile(root, profile, external_prices=None, excluded_codes=None):
    log = []
    external_prices = external_prices or {}
    excluded_codes = excluded_codes or set()
    item_tag = profile["item_tag"]
    items_src = root.xpath(f"//*[local-name()='{item_tag}']") if item_tag else []
    shop = etree.Element("SHOP")
    if not item_tag:
        return shop, ["⚠️ Ještě není vybraný element produktu (krok 2 nebo 3)."]

    log.append(f"Nalezeno {len(items_src)} produktů podle elementu <{item_tag}>.")

    n_price_missing = n_img = n_marketing = n_technical = n_size = n_video = n_video_skipped = 0

    code_src_field = profile["fields"].get("CODE", {}).get("source")
    for src in items_src:
        if excluded_codes and code_src_field and (gtext_deep(src, code_src_field) or "") in excluded_codes:
            continue  # produkt vyřazený ve výběru - vůbec se nezpracovává
        item = etree.SubElement(shop, "SHOPITEM")

        for target, fcfg in profile["fields"].items():
            if fcfg.get("source"):
                val = gtext_deep(src, fcfg["source"])
                if val:
                    etree.SubElement(item, target).text = val

        code_val = gtext(item, "CODE")

        name_el = item.xpath("./*[local-name()='NAME']")
        if name_el and profile.get("name_cleanup", {}).get("enabled"):
            name_el[0].text = clean_name_text(name_el[0].text, code_val, profile["name_cleanup"])

        # cena
        pcfg = profile["price"]
        price_val = None
        if pcfg["mode"] == "direct" and pcfg["source"]:
            price_val = gtext_deep(src, pcfg["source"])
        elif pcfg["mode"] == "price_type" and pcfg["source"]:
            cands = src.xpath(
                f".//*[local-name()='ARTICLE_PRICE'][@{pcfg['price_type_attr']}='{pcfg['price_type_value']}']"
            ) if pcfg.get("price_type_attr") else src.xpath(".//*[local-name()='ARTICLE_PRICE']")
            if not cands:
                cands = src.xpath(".//*[local-name()='ARTICLE_PRICE']")
            if cands:
                price_val = gtext_deep(cands[0], pcfg["source"])
        elif pcfg["mode"] == "external":
            match_key = pcfg.get("match_key", "CODE")
            ident = code_val if match_key == "CODE" else gtext(item, "EAN")
            if ident:
                price_val = external_prices.get(normalize_id(ident))
        if price_val:
            etree.SubElement(item, "CURRENCY").text = profile["currency"]
            try:
                etree.SubElement(item, "PRICE").text = f"{float(price_val.replace(',', '.')):.2f}"
            except ValueError:
                etree.SubElement(item, "PRICE").text = price_val
        else:
            n_price_missing += 1

        # kategorie
        ccfg = profile["category"]
        cats_el = etree.SubElement(item, "CATEGORIES")
        if ccfg["mode"] == "fixed":
            etree.SubElement(cats_el, "CATEGORY").text = ccfg["fixed_value"]
        elif ccfg["mode"] == "direct" and ccfg["source"]:
            for e in gall_deep(src, ccfg["source"]):
                if e.text and e.text.strip():
                    etree.SubElement(cats_el, "CATEGORY").text = e.text.strip()
        if len(cats_el) == 0:
            item.remove(cats_el)

        # váha a rozměry -> LOGISTIC
        wcfg = profile["weight"]
        w = None
        if wcfg["mode"] == "direct" and wcfg["source"]:
            w = gtext_deep(src, wcfg["source"])
        elif wcfg["mode"] == "feature":
            names = set(n.strip().lower() for n in wcfg["names"].split(","))
            for feat in src.xpath(f".//*[local-name()='{wcfg['feature_element']}']"):
                fname = gtext(feat, wcfg["fname_field"])
                if fname and fname.strip().lower() in names:
                    val = gtext(feat, wcfg["fvalue_field"])
                    unit = (gtext(feat, wcfg["funit_field"]) or "kg").lower() if wcfg["funit_field"] else "kg"
                    if val:
                        try:
                            v = float(val.replace(",", "."))
                            if unit == "g":
                                v /= 1000.0
                            w = str(round(v, 3))
                        except ValueError:
                            w = val
                    break

        dims = {}
        if wcfg["mode"] == "feature":
            # Shoptet pro rozměry používá DEPTH / WIDTH / HEIGHT (element LENGTH není povolen),
            # hodnoty v cm - proto se mm/m z feedu přepočítají
            for target, names_key in (("DEPTH", "length_names"), ("WIDTH", "width_names"), ("HEIGHT", "height_names")):
                names_raw = wcfg.get(names_key, "")
                if not names_raw:
                    continue
                dim_names = set(n.strip().lower() for n in names_raw.split(",") if n.strip())
                for feat in src.xpath(f".//*[local-name()='{wcfg['feature_element']}']"):
                    fname = gtext(feat, wcfg["fname_field"])
                    if fname and fname.strip().lower() in dim_names:
                        val = gtext(feat, wcfg["fvalue_field"])
                        dunit = (gtext(feat, wcfg["funit_field"]) or "").strip().lower() if wcfg.get("funit_field") else ""
                        if val:
                            try:
                                dv = float(val.replace(",", "."))
                                if dunit == "mm":
                                    dv /= 10.0
                                elif dunit == "m":
                                    dv *= 100.0
                                dims[target] = ("%.2f" % dv).rstrip("0").rstrip(".")
                            except ValueError:
                                dims[target] = val
                        break

        if w or dims:
            logistic = etree.SubElement(item, "LOGISTIC")
            for target in ("DEPTH", "WIDTH", "HEIGHT"):
                if target in dims:
                    etree.SubElement(logistic, target).text = dims[target]
            if w:
                etree.SubElement(logistic, "WEIGHT").text = w

        # obrázky
        icfg = profile["images"]
        images = []
        if icfg["mode"] == "direct" and icfg["source"]:
            images = [e.text.strip() for e in gall_deep(src, icfg["source"]) if e.text and e.text.strip()]
        elif icfg["mode"] == "mime":
            seen = set()
            for mime in src.xpath(f".//*[local-name()='{icfg['mime_element']}']"):
                purpose = gtext(mime, icfg["purpose_field"]) if icfg["purpose_field"] else None
                if icfg["purpose_field"] and icfg["purpose_value"] and purpose != icfg["purpose_value"]:
                    continue
                fsrc = gtext(mime, icfg["source_field"])
                if fsrc and fsrc not in seen:
                    seen.add(fsrc)
                    images.append(fsrc)
        if icfg.get("max_images"):
            images = images[: icfg["max_images"]]
        base = icfg.get("public_base_url")
        if base:
            base = base if base.endswith("/") else base + "/"
            images = [img if img.lower().startswith(("http://", "https://")) else base + quote(img) for img in images]
        if images:
            n_img += len(images)
            images_el = etree.SubElement(item, "IMAGES")
            item_name = gtext(item, "NAME") or ""
            for i, f in enumerate(images, start=1):
                img_el = etree.SubElement(images_el, "IMAGE")
                img_el.text = f
                if icfg.get("add_description", True) and item_name:
                    desc = item_name if len(images) == 1 or i == 1 else f"{item_name} {i}"
                    img_el.set("description", desc)

        # videa
        vcfg = profile.get("videos", {})
        if vcfg.get("mode") == "mime":
            vids, seen_v = [], set()
            for mime in src.xpath(f".//*[local-name()='{vcfg.get('mime_element', 'MIME')}']"):
                type_field = vcfg.get("type_field")
                mtype = gtext(mime, type_field) if type_field else None
                descr_field = vcfg.get("descr_field")
                descr = (gtext(mime, descr_field) or "") if descr_field else ""
                src_val = gtext(mime, vcfg.get("source_field", "MIME_SOURCE"))
                if not src_val:
                    continue
                type_ok = (not vcfg.get("type_value")) or (mtype == vcfg["type_value"])
                descr_ok = (not vcfg.get("descr_match")) or (vcfg["descr_match"].lower() in descr.lower())
                looks_like_video = any(h in src_val.lower() for h in ("youtube.com", "youtu.be", "vimeo.com"))
                if type_ok and (descr_ok or looks_like_video):
                    yid = extract_youtube_id(src_val)
                    if not yid:
                        n_video_skipped += 1  # ne-YouTube odkaz - Shoptet umí jen YOUTUBE_VIDEO_CODE
                        continue
                    if yid not in seen_v:
                        seen_v.add(yid)
                        vids.append(yid)
            if vids:
                wrapper = vcfg.get("wrapper_tag") or "RELATED_VIDEOS"
                if wrapper == "VIDEOS":  # starý výchozí název z dřívějších šablon
                    wrapper = "RELATED_VIDEOS"
                wrap = etree.SubElement(item, wrapper)
                vtext = gtext(item, "NAME") or ""
                for yid in vids:
                    rv = etree.SubElement(wrap, "RELATED_VIDEO")
                    if vcfg.get("add_text") and vtext:
                        etree.SubElement(rv, "TEXT").text = vtext
                    etree.SubElement(rv, "YOUTUBE_VIDEO_CODE").text = yid
                n_video += len(vids)

        # související / podobné produkty (ARTICLE_REFERENCE apod.)
        rcfg = profile.get("references", {})
        if rcfg.get("enabled") and rcfg.get("included_types"):
            codes_out, seen_r = [], set()
            for ref in src.xpath(f".//*[local-name()='{rcfg.get('reference_element', 'ARTICLE_REFERENCE')}']"):
                rtype = ref.get(rcfg.get("type_attr", "type"))
                if rtype not in rcfg["included_types"]:
                    continue
                idval = gtext(ref, rcfg.get("id_field", "ART_ID_TO"))
                if idval and idval not in seen_r:
                    seen_r.add(idval)
                    codes_out.append(idval)
            if codes_out:
                wrap = etree.SubElement(item, rcfg.get("wrapper_tag", "RELATED_PRODUCTS"))
                for c in codes_out:
                    etree.SubElement(wrap, rcfg.get("item_tag", "CODE")).text = c

        # parametry - klasifikace podle hodnoty vybraného pole
        pcfg2 = profile["params"]
        marketing_bullets, technical_pairs, size_pairs, packaging_bullets = [], [], [], []
        if pcfg2["repeat_element"] and pcfg2["fname_field"] and pcfg2["fvalue_field"]:
            for feat in src.xpath(f".//*[local-name()='{pcfg2['repeat_element']}']"):
                cls_val = None
                if pcfg2.get("classifier_field"):
                    if pcfg2.get("classifier_scope") == "parent":
                        parent = feat.getparent()
                        cls_val = gtext(parent, pcfg2["classifier_field"]) if parent is not None else None
                    else:
                        cls_val = gtext(feat, pcfg2["classifier_field"])
                bucket = pcfg2["value_map"].get(cls_val, "ignore")
                fname = gtext(feat, pcfg2["fname_field"])
                fvals = [e.text.strip() for e in feat.xpath(f"./*[local-name()='{pcfg2['fvalue_field']}']") if e.text and e.text.strip()]
                funit = gtext(feat, pcfg2["funit_field"]) if pcfg2["funit_field"] else None
                if bucket == "marketing" and (fvals or fname):
                    marketing_bullets.extend(fvals if fvals else [fname])
                elif bucket == "packaging" and (fvals or fname):
                    packaging_bullets.extend(fvals if fvals else [fname])
                elif bucket == "technical" and fname and fvals:
                    technical_pairs.append((fname, "; ".join(fvals), funit))
                elif bucket == "size" and fname and fvals:
                    size_pairs.append((fname, "; ".join(fvals), funit))

        desc_bullets = []
        if marketing_bullets:
            n_marketing += len(marketing_bullets)
            text = " • ".join(marketing_bullets)
            if pcfg2["marketing_target"] == "SHORT_DESCRIPTION":
                existing = item.xpath("./*[local-name()='SHORT_DESCRIPTION']")
                if existing:
                    existing[0].text = (existing[0].text or "") + " • " + text
                else:
                    etree.SubElement(item, "SHORT_DESCRIPTION").text = text
            else:
                desc_bullets = list(marketing_bullets)

        if technical_pairs:
            n_technical += len(technical_pairs)
            params_el = etree.SubElement(item, "INFORMATION_PARAMETERS")
            for fname, fval, funit in technical_pairs:
                p = etree.SubElement(params_el, "INFORMATION_PARAMETER")
                etree.SubElement(p, "NAME").text = fname
                etree.SubElement(p, "VALUE").text = f"{fval} {funit}".strip() if funit else fval

        if size_pairs:
            n_size += len(size_pairs)
            params_el = etree.SubElement(item, "PARAMETERS")
            for fname, fval, funit in size_pairs:
                p = etree.SubElement(params_el, "PARAMETER")
                etree.SubElement(p, "NAME").text = fname
                etree.SubElement(p, "VALUE").text = f"{fval} {funit}".strip() if funit else fval

        # DESCRIPTION: popis + odrážky + (úplně na konci) obsah balení, s kontrolou duplicit
        dcfg = profile.get("description_format", {})
        desc_el = item.xpath("./*[local-name()='DESCRIPTION']")
        raw_desc = desc_el[0].text if desc_el else ""
        if dcfg.get("enabled") or desc_bullets or packaging_bullets:
            new_desc = format_description_html(raw_desc, dcfg, desc_bullets, packaging_bullets)
            if new_desc:
                if desc_el:
                    desc_el[0].text = new_desc
                else:
                    etree.SubElement(item, "DESCRIPTION").text = new_desc

        etree.SubElement(item, "VISIBILITY").text = profile["visibility_default"]

    if n_price_missing:
        log.append(f"⚠️ {n_price_missing} produktů nemá dohledatelnou cenu.")
    if n_img:
        log.append(f"Nalezeno {n_img} obrázkových položek.")
    if n_video:
        log.append(f"{n_video} YouTube videí zapsáno do <RELATED_VIDEOS>.")
    if n_video_skipped:
        log.append(f"⚠️ {n_video_skipped} video odkazů přeskočeno - nejsou z YouTube (Shoptet umí jen YOUTUBE_VIDEO_CODE).")
    if n_marketing:
        log.append(f"{n_marketing} marketingových textů zapsáno do popisu.")
    if n_technical:
        log.append(f"{n_technical} technických parametrů do <INFORMATION_PARAMETERS>.")
    if n_size:
        log.append(f"{n_size} velikostních parametrů do <PARAMETERS>.")
    apply_text_tools(shop, profile)
    return shop, log


def apply_name_affixes(shop_root, prefix, suffix):
    if not prefix and not suffix:
        return shop_root
    for item in shop_root.xpath("./*[local-name()='SHOPITEM']"):
        name_el = item.xpath("./*[local-name()='NAME']")
        if name_el and name_el[0].text:
            parts = [p for p in [(prefix or "").strip(), name_el[0].text.strip(), (suffix or "").strip()] if p]
            name_el[0].text = " ".join(parts)
    return shop_root


OUT_OF_STOCK_OPTIONS = [
    "2-3 dny",
    "Obvykle do 7 dní",
    "Produkt již není v prodeji",
    "Předprodej",
    "Na dotaz",
]


def apply_availability(shop_root, cfg):
    """Hromadně nastaví text dostupnosti skladem/nedostupné pro všechny produkty."""
    if not cfg or not cfg.get("enabled"):
        return shop_root
    in_stock = (cfg.get("in_stock") or "").strip()
    out_of_stock = (cfg.get("out_of_stock") or "").strip()
    for item in shop_root.xpath("./*[local-name()='SHOPITEM']"):
        if in_stock:
            el = item.xpath("./*[local-name()='AVAILABILITY_IN_STOCK']")
            (el[0] if el else etree.SubElement(item, "AVAILABILITY_IN_STOCK")).text = in_stock
        if out_of_stock:
            el = item.xpath("./*[local-name()='AVAILABILITY_OUT_OF_STOCK']")
            (el[0] if el else etree.SubElement(item, "AVAILABILITY_OUT_OF_STOCK")).text = out_of_stock
    return shop_root


def best_net_for_gross(gross, vat_percent):
    """Cena BEZ DPH na max. 2 desetinná místa, jejíž cena s DPH (po Shoptetově přičtení) je
    nejblíž cílové částce 'gross'. Zkouší i sousední haléře, protože prosté zaokrouhlení
    G/(1+DPH) na 2 místa občas vyjde o 0,01 vedle."""
    f = 1 + vat_percent / 100.0
    base = round(gross / f, 2)
    cands = [round(base + d / 100.0, 2) for d in (-1, 0, 1)]
    return min(cands, key=lambda n: (abs(round(n * f, 2) - gross), abs(n * f - gross)))


def apply_price_precision(shop_root, mode):
    """Finální průchod při exportu: PRICE a STANDARD_PRICE mají nejvýš 2 desetinná místa
    ('2dp'), nebo se zaokrouhlí na celá čísla ('int')."""
    for item in shop_root.xpath("./*[local-name()='SHOPITEM']"):
        for tag in ("PRICE", "STANDARD_PRICE"):
            el = item.xpath(f"./*[local-name()='{tag}']")
            if el and el[0].text:
                try:
                    v = float(el[0].text.replace(",", "."))
                except ValueError:
                    continue
                el[0].text = str(int(v + 0.5)) if mode == "int" else f"{v:.2f}"
    return shop_root


def apply_price_adjustment(shop_root, cfg):
    """Hromadná úprava ceny.
    Cílem je vždy CELÉ číslo jako výsledná zobrazená cena v Shoptetu. Shoptet si k <PRICE>
    DPH podle <VAT> vždy dopočítává sám (bere <PRICE> jako cenu bez DPH), takže nejde do
    <PRICE> poslat rovnou kulatou cenu s DPH - to by Shoptet vzal jako základ a k tomu
    ještě připočetl DPH podruhé. Proto se to počítá "pozpátku": nejdřív se spočte cílová
    KULATÁ cena s DPH (podle toho, jestli zdrojová cena z feedu už DPH obsahovala, nebo ne),
    a do <PRICE> (resp. <STANDARD_PRICE>) se zapíše taková cena BEZ DPH, aby po Shoptetově
    vlastním přičtení DPH vyšla přesně ta kulatá částka.
    Čistá funkce na hotovém stromu, takže funguje nezávisle na tom, odkud cena původně přišla."""
    if not cfg or not cfg.get("enabled"):
        return shop_root
    discount = cfg.get("discount_percent") or 0
    source_mode = cfg.get("source_mode", "net")  # "net" = zdrojová cena bez DPH, "gross" = zdrojová cena s DPH (např. z ceníku)
    try:
        vat = float(str(cfg.get("vat_percent", "21")).replace(",", "."))
    except (TypeError, ValueError):
        vat = 0
    for item in shop_root.xpath("./*[local-name()='SHOPITEM']"):
        price_el = item.xpath("./*[local-name()='PRICE']")
        if not price_el or not price_el[0].text:
            continue
        try:
            p = float(price_el[0].text.replace(",", "."))
        except ValueError:
            continue

        # cílová (kulatá) zobrazená cena s DPH
        target_list = round(p) if source_mode == "gross" else round(p * (1 + vat / 100.0))
        target_final = round(target_list * (1 - discount / 100.0)) if discount else target_list

        # zapisujeme cenu BEZ DPH (max. 2 desetinná místa - víc validátor Shoptetu nepustí), vybranou
        # tak, aby po Shoptetově přičtení DPH vyšla cena co nejblíž cílové kulaté částce
        fed_list = best_net_for_gross(target_list, vat)
        fed_final = best_net_for_gross(target_final, vat)

        def fmt(v):
            return f"{v:.2f}"

        price_el[0].text = fmt(fed_final)
        if discount:
            std_el = item.xpath("./*[local-name()='STANDARD_PRICE']")
            if std_el:
                std_el[0].text = fmt(fed_list)
            else:
                etree.SubElement(item, "STANDARD_PRICE").text = fmt(fed_list)
        vat_el = item.xpath("./*[local-name()='VAT']")
        if vat_el:
            vat_el[0].text = str(cfg.get("vat_percent", "21"))
        else:
            etree.SubElement(item, "VAT").text = str(cfg.get("vat_percent", "21"))
    return shop_root


def apply_overrides(shop_root, overrides):
    """Ruční opravy z kroku 7 - aplikují se navrch čerstvě vygenerovaného stromu,
    takže přežijí návrat na dřívější krok a přepočítání."""
    if not overrides:
        return shop_root
    for item in shop_root.xpath("./*[local-name()='SHOPITEM']"):
        code = gtext(item, "CODE") or ""
        for tag, value in overrides.get(code, {}).items():
            if tag == "CATEGORY":
                cats_el = item.xpath("./*[local-name()='CATEGORIES']")
                if not cats_el:
                    cats_el = [etree.SubElement(item, "CATEGORIES")]
                etree.SubElement(cats_el[0], "CATEGORY").text = value
                continue
            existing = item.xpath(f"./*[local-name()='{tag}']")
            if existing:
                existing[0].text = value
            else:
                etree.SubElement(item, tag).text = value
    return shop_root


def assign_doc_order(root):
    return {el: i for i, el in enumerate(root.iter())}


def validate_shop(shop_root):
    order = assign_doc_order(shop_root)
    issues = []
    for item in shop_root.xpath("./*[local-name()='SHOPITEM']"):
        code = gtext(item, "CODE") or "(bez kódu)"
        name = gtext(item, "NAME") or "(bez názvu)"
        pos = order.get(item, 0)

        def missing(tag, mandatory, message):
            issues.append({"tag": tag, "mandatory": mandatory, "code": code, "name": name, "item": item,
                            "doc_pos": pos, "message": message})

        if not gtext(item, "CODE"):
            missing("CODE", True, "Chybí povinný CODE.")
        if not gtext(item, "NAME"):
            missing("NAME", True, "Chybí povinný NAME.")
        if not gtext(item, "PRICE"):
            missing("PRICE", True, "Chybí povinná cena PRICE.")
        if not gtext(item, "CURRENCY"):
            missing("CURRENCY", True, "Chybí povinná měna CURRENCY.")
        if not item.xpath("./*[local-name()='CATEGORIES']/*[local-name()='CATEGORY']"):
            missing("CATEGORY", True, "Produkt nemá žádnou kategorii.")
        if not gtext(item, "MANUFACTURER"):
            missing("MANUFACTURER", False, "Doporučeno vyplnit výrobce.")
        if not gtext(item, "EAN"):
            missing("EAN", False, "Doporučeno vyplnit EAN.")
        if not item.xpath("./*[local-name()='IMAGES']/*[local-name()='IMAGE']"):
            missing("IMAGES", False, "Produkt nemá žádný obrázek.")
        else:
            for img in item.xpath("./*[local-name()='IMAGES']/*[local-name()='IMAGE']"):
                txt = (img.text or "").strip()
                if txt and not txt.lower().startswith(("http://", "https://")):
                    issues.append({"tag": "IMAGE_URL", "mandatory": True, "code": code, "name": name, "item": item,
                                   "doc_pos": order.get(img, pos), "message": f"IMAGE '{txt}' není platná URL."})
        if not gtext(item, "VISIBILITY"):
            missing("VISIBILITY", False, "Chybí VISIBILITY.")
    issues.sort(key=lambda x: (not x["mandatory"], -x["doc_pos"]))
    return issues


def read_table_file(uploaded_file):
    """Přečte Excel/CSV do (header, rows) - header je list názvů sloupců (nebo indexy jako
    text, pokud soubor hlavičku nemá), rows je list řádků (list hodnot jako string)."""
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    rows = []
    if name.endswith(".csv"):
        import csv as csv_mod
        import io as io_mod
        text = data.decode("utf-8", errors="ignore")
        for row in csv_mod.reader(io_mod.StringIO(text)):
            rows.append([c.strip() for c in row])
    elif name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        import io as io_mod
        wb = openpyxl.load_workbook(io_mod.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c).strip() for c in row])
    if not rows:
        return [], []
    header = rows[0]
    return header, rows[1:]


def normalize_id(v):
    v = str(v).strip()
    if v.endswith(".0") and v[:-2].replace("-", "").isdigit():
        v = v[:-2]
    return v


def read_all_columns(uploaded_file):
    """Přečte nahraný soubor do slovníku {název_listu: list sloupců}, každý sloupec = list hodnot (str).
    Excel se čte bez read_only (ten umí useknout řádky podle chybné dimenze listu)."""
    import io as io_mod
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read()

    def cell_str(c):
        if c is None:
            return ""
        if isinstance(c, float) and c.is_integer():
            return str(int(c))
        return str(c)

    sheets = {}
    if name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="ignore")
        vals = [p for p in re.split(r"[\r\n,;\t]+", text)]
        sheets["TXT"] = [vals]
    elif name.endswith(".csv"):
        import csv as csv_mod
        text = data.decode("utf-8-sig", errors="ignore")
        try:
            dialect = csv_mod.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except Exception:
            dialect = csv_mod.excel
        rows = list(csv_mod.reader(io_mod.StringIO(text), dialect))
        w = max((len(r) for r in rows), default=0)
        sheets["CSV"] = [[(r[i] if i < len(r) else "") for r in rows] for i in range(w)]
    elif name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io_mod.BytesIO(data), data_only=True)
        for ws in wb.worksheets:
            rows = [[cell_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
            w = max((len(r) for r in rows), default=0)
            if w:
                sheets[ws.title] = [[(r[i] if i < len(r) else "") for r in rows] for i in range(w)]
    elif name.endswith(".xls"):
        df = pd.read_excel(io_mod.BytesIO(data), header=None, dtype=str)
        sheets["List1"] = [df[c].fillna("").tolist() for c in df.columns]
    return sheets


def norm_code(v, loose=False):
    """Normalizace kódu pro porovnání. Přísná: ořez mezer (vč. nbsp), '.0' u čísel, velká písmena.
    Volná: navíc bez oddělovačů (mezera . - _ /) a bez úvodních nul (Excel je u čísel ztrácí)."""
    s = str(v if v is not None else "")
    s = s.replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "").strip()
    s = re.sub(r"^(\d+)\.0+$", r"\1", s)
    s = s.upper()
    if loose:
        s = re.sub(r"[\s.\-_/]+", "", s).lstrip("0")
    return s


def match_codes(upload_values, feed_pairs):
    """upload_values: list hodnot z Excelu. feed_pairs: list (hodnota_pole_ve_feedu, CODE).
    Vrací (set CODE, počet přesných, počet volných, list nespárovaných hodnot)."""
    strict = {}
    loose = {}
    for val, code in feed_pairs:
        if not val:
            continue
        strict.setdefault(norm_code(val), set()).add(code)
        loose.setdefault(norm_code(val, True), set()).add(code)
    matched, n_strict, n_loose, unmatched = set(), 0, 0, []
    seen = set()
    for v in upload_values:
        k = norm_code(v)
        if not k or k in seen:
            continue
        seen.add(k)
        if k in strict:
            matched |= strict[k]
            n_strict += 1
            continue
        lk = norm_code(v, True)
        if lk and lk in loose:
            matched |= loose[lk]
            n_loose += 1
        else:
            unmatched.append(str(v).strip())
    return matched, n_strict, n_loose, unmatched


def shop_to_bytes(shop_root):
    return etree.tostring(shop_root, xml_declaration=True, encoding="utf-8", pretty_print=True)


# =============================================================================
# SESSION STATE
# =============================================================================
ss = st.session_state
ss.setdefault("raw_root", None)
ss.setdefault("profile", json.loads(json.dumps(DEFAULT_PROFILE)))
ss.setdefault("overrides", {})
ss.setdefault("excluded_codes", set())
ss.setdefault("step", 1)
ss.setdefault("external_prices", {})


def get_item_candidates(root):
    """Kandidáty na element produktu - výpočet je u velkých feedů pomalý, proto se cachuje."""
    cached = ss.get("_item_candidates")
    if cached is not None and cached[0] == ss.get("_raw_ver", 0):
        return cached[1]
    res = candidate_item_tags(root)
    ss._item_candidates = (ss.get("_raw_ver", 0), res)
    return res


def compute_shop():
    """Vrací (strom, log). Výsledek se cachuje podle nastavení - volající ho nesmí měnit."""
    if ss.raw_root is None:
        return None, []
    key = json.dumps([ss.profile, ss.overrides, sorted(ss.excluded_codes),
                      ss.get("_raw_ver", 0), ss.get("_ext_ver", 0)], sort_keys=True, default=str)
    cached = ss.get("_shop_cache")
    if cached and cached[0] == key:
        return cached[1], cached[2]
    if detect_feed_type(ss.raw_root) == "shoptet":
        shop = deepcopy(ss.raw_root)
        for it in shop.xpath("./*[local-name()='SHOPITEM']"):
            if (gtext(it, "CODE") or "") in ss.excluded_codes:
                shop.remove(it)
        log = ["Feed byl už ve formátu Shoptet, načten beze změny."]
    else:
        shop, log = apply_profile(ss.raw_root, ss.profile, ss.get("external_prices", {}), ss.excluded_codes)
    shop = apply_name_affixes(shop, ss.profile.get("name_prefix", ""), ss.profile.get("name_suffix", ""))
    shop = apply_price_adjustment(shop, ss.profile.get("price_adjust", {}))
    shop = apply_availability(shop, ss.profile.get("availability", {}))
    shop = apply_overrides(shop, ss.overrides)
    shop = apply_price_precision(shop, ss.profile.get("price_precision", "2dp"))
    ss._shop_cache = (key, shop, log)
    return shop, log


def render_live_preview():
    shop, _ = compute_shop()
    if shop is None:
        return
    issues = validate_shop(shop)
    n_mand = sum(1 for i in issues if i["mandatory"])
    n_opt = len(issues) - n_mand
    items = shop.xpath("./*[local-name()='SHOPITEM']")
    c1, c2, c3 = st.columns(3)
    c1.metric("Produktů", len(items))
    c2.metric("🔴 Povinné chyby", n_mand)
    c3.metric("⚪ Doporučené", n_opt)
    if n_mand:
        st.error(f"Feed zatím NENÍ validní pro Shoptet - {n_mand} povinných chyb (viz krok 7).")
    elif items:
        st.success("Feed prochází kontrolou povinných elementů.")
    with st.expander("👁️ Živý náhled XML (aktuální stav)", expanded=False):
        xml_bytes = shop_to_bytes(shop)
        st.code(xml_bytes.decode("utf-8")[:8000], language="xml")
        if len(xml_bytes) > 8000:
            st.caption("(zkráceno, plný výstup je v kroku 8 - Export)")


def nav_buttons():
    c1, c2, c3 = st.columns([1, 3, 1])
    with c1:
        if ss.step > 1 and st.button("⬅️ Zpět"):
            ss.step -= 1
            st.rerun()
    with c3:
        if ss.step < len(STEPS) and st.button("Další ➡️", type="primary"):
            ss.step += 1
            st.rerun()


# =============================================================================
# HLAVIČKA
# =============================================================================
st.title("🛠️ Univerzální Feed → Shoptet Validator")
progress_labels = " → ".join(f"**{s}**" if i + 1 == ss.step else s for i, s in enumerate(STEPS))
st.caption(progress_labels)
step_jump = st.selectbox("Přeskočit na krok", STEPS, index=ss.step - 1, label_visibility="collapsed")
if STEPS.index(step_jump) + 1 != ss.step:
    ss.step = STEPS.index(step_jump) + 1
    st.rerun()
st.markdown("---")

with st.sidebar:
    st.header("📖 Slovník Shoptet tagů")
    search_term = st.text_input("Hledat tag / popis:", "").lower()
    for tag, info in SHOPTET_DICT.items():
        if search_term in tag.lower() or search_term in info["desc"].lower():
            badge = "🔴 POVINNÝ" if info["mandatory"] else "⚪ nepovinný"
            st.markdown(f"**`<{tag}>`** — {badge}")
            st.caption(info["desc"])

step = ss.step

# =============================================================================
# KROK 1 - NAHRÁNÍ
# =============================================================================
if step == 1:
    st.header(STEPS[0])
    up_col1, up_col2 = st.columns(2)
    with up_col1:
        uploaded_file = st.file_uploader("Nahraj vstupní XML feed", type=["xml"])
    with up_col2:
        profile_file = st.file_uploader("Nebo nahraj uloženou šablonu (.json) a pak i feed", type=["json"])

    if uploaded_file is not None:
        data = uploaded_file.read()
        if ss.get("_last_filename") != uploaded_file.name:
            ss.raw_root = parse_xml_bytes(data)
            ss._last_filename = uploaded_file.name
            ss._raw_ver = ss.get("_raw_ver", 0) + 1
            ss.overrides = {}
            ss.excluded_codes = set()

    if profile_file is not None and not ss.get("_profile_loaded"):
        try:
            ss.profile = json.load(profile_file)
            ss._profile_loaded = True
            st.success(f"Šablona '{ss.profile.get('profile_name', '')}' načtena.")
        except Exception as e:
            st.error(f"Nepodařilo se načíst šablonu: {e}")

    if ss.raw_root is not None:
        st.success("Feed nahrán.")
        st.caption(f"Typ: {detect_feed_type(ss.raw_root)}")
    nav_buttons()

elif ss.raw_root is None:
    st.warning("Nejdřív nahraj feed v kroku 1.")
    if st.button("⬅️ Jít na krok 1"):
        ss.step = 1
        st.rerun()

# =============================================================================
# KROK 2 - VÝBĚR PRODUKTŮ (hned po nahrání - nevybrané se dál vůbec nezpracovávají)
# =============================================================================
elif step == 2:
    st.header(STEPS[1])
    root = ss.raw_root
    prof = ss.profile
    st.caption("Vyber produkty hned na začátku - nezaškrtnuté se dál vůbec nezpracovávají, takže další kroky "
               "pracují s menším množstvím dat a jsou rychlejší.")

    rows_src = []
    items_src = []
    tag_options = []
    ready = True
    if detect_feed_type(root) == "shoptet":
        for it in root.xpath("./*[local-name()='SHOPITEM']"):
            rows_src.append((gtext(it, "CODE") or "", gtext(it, "NAME") or ""))
            items_src.append(it)
    else:
        candidates = get_item_candidates(root)
        cand_names = [f"{t} ({c}×)" for t, c in candidates]
        cand_map = {f"{t} ({c}×)": t for t, c in candidates}
        default_idx = next((i for i, (t, c) in enumerate(candidates) if t == prof.get("item_tag")), 0)
        chosen = st.selectbox("Element představující JEDEN produkt", cand_names,
                               index=default_idx if cand_names else 0)
        prof["item_tag"] = cand_map.get(chosen, prof.get("item_tag", ""))
        sample_items = root.xpath(f"//*[local-name()='{prof['item_tag']}']") if prof["item_tag"] else []
        avail_tags = aggregate_descendant_tags(sample_items) if sample_items else []
        tag_options = ["-- nevybráno --"] + avail_tags
        c1, c2 = st.columns(2)
        with c1:
            cur = prof["fields"].get("CODE", {"source": ""}).get("source", "")
            sel = st.selectbox("Pole s kódem produktu (CODE)", tag_options,
                                index=tag_options.index(cur) if cur in tag_options else 0, key="f_CODE")
            prof["fields"]["CODE"] = {"source": "" if sel == "-- nevybráno --" else sel}
        with c2:
            cur = prof["fields"].get("NAME", {"source": ""}).get("source", "")
            sel = st.selectbox("Pole s názvem produktu (NAME)", tag_options,
                                index=tag_options.index(cur) if cur in tag_options else 0, key="f_NAME")
            prof["fields"]["NAME"] = {"source": "" if sel == "-- nevybráno --" else sel}
        code_src = prof["fields"]["CODE"]["source"]
        name_src = prof["fields"]["NAME"]["source"]
        if not code_src:
            ready = False
            st.info("Vyber pole s kódem produktu - podle něj se produkty vybírají.")
        else:
            for it in sample_items:
                rows_src.append((gtext_deep(it, code_src) or "", gtext_deep(it, name_src) or ""))
                items_src.append(it)

    if ready:
        all_codes = set(c for c, _ in rows_src)
        excluded_now = ss.excluded_codes & all_codes
        df_sel = pd.DataFrame([{"Zahrnout": code not in excluded_now, "CODE": code, "NAME": name}
                                for code, name in rows_src])
        c1, c2, c3 = st.columns(3)
        with c1:
            if st.button("Zaškrtnout vše"):
                ss.excluded_codes = set()
                st.rerun()
        with c2:
            if st.button("Odškrtnout vše"):
                ss.excluded_codes = set(all_codes)
                st.rerun()
        with c3:
            st.metric("Vybráno", f"{len(rows_src) - len(excluded_now)} / {len(rows_src)}")
        edited = st.data_editor(df_sel, use_container_width=True, disabled=["CODE", "NAME"], hide_index=True, key="sel_editor")
        if st.button("💾 Uložit výběr", type="primary"):
            ss.excluded_codes = set(edited.loc[~edited["Zahrnout"], "CODE"])
            st.rerun()

        st.markdown("---")
        st.subheader("📄 Nebo nahraj seznam kódů ze souboru (Excel/CSV/TXT)")
        st.caption("Vyber sloupec s kódy - porovná se s polem feedu, které zvolíš níže (výchozí CODE). "
                   "Porovnání ignoruje mezery, velikost písmen a koncovku .0; pokud nenajde přesnou shodu, zkusí i "
                   "variantu bez oddělovačů a úvodních nul (Excel je u čísel ztrácí).")
        code_file = st.file_uploader("Soubor se seznamem kódů", type=["xlsx", "xlsm", "xls", "csv", "txt"],
                                     key="code_file_uploader")
        if code_file is not None:
            try:
                sheets = read_all_columns(code_file)
            except Exception as e:
                sheets = {}
                st.error(f"Soubor se nepodařilo přečíst: {e}")
            if sheets:
                cs1, cs2, cs3 = st.columns(3)
                with cs1:
                    sheet_name = st.selectbox("List", list(sheets.keys()))
                cols = sheets[sheet_name]
                nonempty = [i for i, c in enumerate(cols) if any(str(x).strip() for x in c)]
                with cs2:
                    col_idx = st.selectbox(
                        "Sloupec s kódy", nonempty or [0],
                        format_func=lambda i: f"Sloupec {i + 1}  (např. {next((str(x).strip() for x in cols[i] if str(x).strip()), '')[:20]}, "
                                              f"{sum(1 for x in cols[i] if str(x).strip())} hodnot)" if cols else "-")
                feed_fields = [t for t in (tag_options[1:] if tag_options else [])]
                default_field = prof["fields"].get("CODE", {}).get("source", "") if prof.get("fields") else ""
                with cs3:
                    if feed_fields:
                        cmp_field = st.selectbox("Porovnat s polem feedu", feed_fields,
                                                 index=feed_fields.index(default_field) if default_field in feed_fields else 0)
                    else:
                        cmp_field = None
                        st.text_input("Porovnat s polem feedu", "CODE", disabled=True)
                values = list(cols[col_idx]) if cols else []
                feed_pairs = []
                for it, (code, _n) in zip(items_src, rows_src):
                    fv = gtext_deep(it, cmp_field) if cmp_field else code
                    feed_pairs.append((fv or "", code))
                # hlavička: první hodnota, která nic nenapáruje, a vypadá jako text (ne číslo)
                first = str(values[0]).strip() if values else ""
                m_first = match_codes([first], feed_pairs)[0] if first else set()
                has_header = st.checkbox("První řádek je hlavička", value=bool(first) and not m_first and not first.isdigit())
                if has_header:
                    values = values[1:]
                values = [v for v in values if str(v).strip()]
                matched, n_strict, n_loose, unmatched = match_codes(values, feed_pairs)
                n_unique = n_strict + n_loose + len(unmatched)
                st.success(f"V souboru **{n_unique}** unikátních kódů → **{n_strict + n_loose}** nalezeno ve feedu "
                           f"({n_strict} přesně, {n_loose} po volné normalizaci), **{len(unmatched)}** nenalezeno. "
                           f"Ve feedu je celkem {len(all_codes)} produktů, vybráno by jich bylo {len(matched)}.")
                if unmatched:
                    with st.expander(f"Nenalezené kódy ({len(unmatched)}) - nejsou ve feedu"):
                        st.write(", ".join(unmatched[:300]) + (" …" if len(unmatched) > 300 else ""))
                        st.caption("Zkontroluj, zda tyto kódy ve feedu opravdu existují (feed dodavatele nemusí obsahovat "
                                   "všechny položky), nebo zkus v rozbalovátku výše porovnat s jiným polem (např. EAN).")
                        sample_feed = [v for v, _c in feed_pairs[:5]]
                        st.caption(f"Ukázka hodnot ve feedu: {', '.join(sample_feed)}")
                cbtn1, cbtn2 = st.columns(2)
                with cbtn1:
                    if st.button("✅ Zaškrtnout JEN tyto kódy", type="primary"):
                        ss.excluded_codes = all_codes - matched
                        st.rerun()
                with cbtn2:
                    if st.button("🚫 Odškrtnout tyto kódy (vynechat je)"):
                        ss.excluded_codes = ss.excluded_codes | matched
                        st.rerun()
    nav_buttons()

# =============================================================================
# KROK 3 - ZÁKLADNÍ MAPOVÁNÍ
# =============================================================================
elif step == 3:
    st.header(STEPS[2])
    root = ss.raw_root
    if detect_feed_type(root) == "shoptet":
        st.success("Tenhle feed už je ve formátu Shoptet - mapování polí není potřeba, pokračuj dál.")
    else:
        prof = ss.profile
        prof["profile_name"] = st.text_input("Název šablony", value=prof.get("profile_name", "Nová šablona"))

        candidates = get_item_candidates(root)
        cand_names = [f"{t} ({c}×)" for t, c in candidates]
        cand_map = {f"{t} ({c}×)": t for t, c in candidates}
        default_idx = next((i for i, (t, c) in enumerate(candidates) if t == prof.get("item_tag")), 0)
        chosen = st.selectbox("Element představující JEDEN produkt", cand_names,
                               index=default_idx if cand_names else 0)
        prof["item_tag"] = cand_map.get(chosen, prof.get("item_tag", ""))

        sample_items = root.xpath(f"//*[local-name()='{prof['item_tag']}']") if prof["item_tag"] else []
        avail_tags = aggregate_descendant_tags(sample_items) if sample_items else []
        tag_options = ["-- nevybráno --"] + avail_tags

        def tag_select(label, current, key):
            idx = tag_options.index(current) if current in tag_options else 0
            return st.selectbox(label, tag_options, index=idx, key=key)

        st.markdown("#### Jednoduchá pole")
        cols = st.columns(3)
        for i, target in enumerate(["CODE", "NAME", "DESCRIPTION", "MANUFACTURER", "EAN", "PART_NUMBER"]):
            with cols[i % 3]:
                cur = prof["fields"].get(target, {"source": ""})["source"]
                sel = tag_select(target, cur, f"f_{target}")
                prof["fields"][target] = {"source": "" if sel == "-- nevybráno --" else sel}

        with st.expander("📝 Formátování DESCRIPTION (z jednoho bloku textu udělat strukturu)", expanded=False):
            dfcfg = prof.get("description_format", {})
            dfcfg["enabled"] = st.checkbox("Zapnout formátování popisu", value=dfcfg.get("enabled", False))
            if dfcfg["enabled"]:
                c1, c2 = st.columns(2)
                with c1:
                    dfcfg["paragraphs"] = st.checkbox("Rozdělit do odstavců", value=dfcfg.get("paragraphs", True))
                    dfcfg["headings"] = st.checkbox("Přidat nadpis 'Popis'", value=dfcfg.get("headings", True))
                with c2:
                    dfcfg["extract_packaging"] = st.checkbox("Vytáhnout obsah balení do vlastní sekce",
                                                               value=dfcfg.get("extract_packaging", True))
                    dfcfg["bold_numbers"] = st.checkbox("Zvýraznit klíčová čísla tučně (2 000 W, 0,6 kg...)",
                                                          value=dfcfg.get("bold_numbers", True))
                dfcfg["dedupe"] = st.checkbox("Kontrolovat duplicity (odrážky/texty, které už v popisu jsou, vynechat)",
                                               value=dfcfg.get("dedupe", True))
                st.caption("Obsah balení se vždy skládá úplně na konec popisu.")
                prof["description_format"] = dfcfg
                desc_source = prof["fields"].get("DESCRIPTION", {}).get("source")
                if desc_source and sample_items:
                    sample_raw = gtext_deep(sample_items[0], desc_source)
                    if sample_raw:
                        st.caption("Náhled na prvním produktu:")
                        st.markdown(format_description_html(sample_raw, dfcfg), unsafe_allow_html=True)
            else:
                prof["description_format"] = dfcfg

        with st.expander("🧹 Čištění NAME (odstranění kódu/balení z názvu)", expanded=False):
            ncfg = prof.get("name_cleanup", {"enabled": False, "strip_after_semicolon": True, "strip_own_code_parens": True})
            ncfg["enabled"] = st.checkbox("Zapnout čištění názvu", value=ncfg.get("enabled", False))
            if ncfg["enabled"]:
                ncfg["strip_after_semicolon"] = st.checkbox(
                    "Odstranit vše za středníkem (typicky typ balení)", value=ncfg.get("strip_after_semicolon", True))
                ncfg["strip_own_code_parens"] = st.checkbox(
                    "Odstranit kód produktu v závorce, pokud je stejný jako CODE", value=ncfg.get("strip_own_code_parens", True))
                name_source = prof["fields"].get("NAME", {}).get("source")
                if name_source and sample_items:
                    sample_code_src = prof["fields"].get("CODE", {}).get("source")
                    sample_raw_name = gtext_deep(sample_items[0], name_source)
                    sample_code = gtext_deep(sample_items[0], sample_code_src) if sample_code_src else None
                    if sample_raw_name:
                        st.caption("Náhled na prvním produktu:")
                        st.write(f"„{sample_raw_name}" + "“ → „" + clean_name_text(sample_raw_name, sample_code, ncfg) + "“")
            prof["name_cleanup"] = ncfg

        with st.expander("✏️ Úpravy textu (hledat/nahradit, velikost písmen, zkrácení)", expanded=False):
            ttcfg = prof.get("text_tools") or {"enabled": False, "collapse_spaces": True, "name_case": "none",
                                               "short_desc_max": 0, "replace_rules": []}
            ttcfg["enabled"] = st.checkbox("Zapnout úpravy textu", value=ttcfg.get("enabled", False))
            if ttcfg["enabled"]:
                ttcfg["collapse_spaces"] = st.checkbox("Odstranit dvojité mezery a mezery na okrajích",
                                                       value=ttcfg.get("collapse_spaces", True))
                case_opts = {"none": "Beze změny", "first_upper": "První písmeno velké",
                             "sentence": "Jako věta (první velké, zbytek malá)", "upper": "VELKÁ PÍSMENA", "lower": "malá písmena"}
                keys = list(case_opts.keys())
                ttcfg["name_case"] = keys[st.selectbox(
                    "Velikost písmen v NAME", range(len(keys)), format_func=lambda i: case_opts[keys[i]],
                    index=keys.index(ttcfg.get("name_case", "none")) if ttcfg.get("name_case", "none") in keys else 0)]
                ttcfg["short_desc_max"] = int(st.number_input(
                    "Zkrátit SHORT_DESCRIPTION na max. znaků (0 = nezkracovat; odstraní HTML)", min_value=0, max_value=2000,
                    value=int(ttcfg.get("short_desc_max", 0)), step=10))
                st.markdown("**Hledat a nahradit** (prázdné 'Nahradit' = smazat; zaškrtni Regex pro regulární výrazy)")
                rules_df = pd.DataFrame(ttcfg.get("replace_rules") or [],
                                        columns=["field", "find", "replace", "regex"])
                if rules_df.empty:
                    rules_df = pd.DataFrame({"field": pd.Series(dtype="str"), "find": pd.Series(dtype="str"),
                                             "replace": pd.Series(dtype="str"), "regex": pd.Series(dtype="bool")})
                rules_df["regex"] = rules_df["regex"].fillna(False).astype(bool)
                edited_rules = st.data_editor(
                    rules_df, num_rows="dynamic", use_container_width=True, key="text_rules_editor",
                    column_config={
                        "field": st.column_config.SelectboxColumn("Pole", options=["VŠECHNA", "NAME", "SHORT_DESCRIPTION", "DESCRIPTION"], default="VŠECHNA"),
                        "find": st.column_config.TextColumn("Hledat"),
                        "replace": st.column_config.TextColumn("Nahradit"),
                        "regex": st.column_config.CheckboxColumn("Regex", default=False)})
                ttcfg["replace_rules"] = [
                    {"field": r["field"] or "VŠECHNA", "find": r["find"] or "", "replace": r["replace"] or "",
                     "regex": bool(r["regex"])}
                    for r in edited_rules.fillna("").to_dict("records") if (r.get("find") or "")]
                for fld in ("NAME", "SHORT_DESCRIPTION", "DESCRIPTION"):
                    src_f = prof["fields"].get(fld, {}).get("source")
                    if src_f and sample_items:
                        raw = gtext_deep(sample_items[0], src_f)
                        if raw:
                            st.caption(f"Náhled {fld} (první produkt; bez dalšího formátování popisu/čištění názvu):")
                            st.write(f"„{raw[:200]}“ → „{apply_text_rules(raw, fld, ttcfg)[:200]}“")
            prof["text_tools"] = ttcfg

        st.markdown("#### Cena")
        pcfg = prof["price"]
        price_mode = st.radio("Zdroj ceny", ["Přímé pole", "BMEcat styl (ARTICLE_PRICE dle typu)", "Ceník ze souboru (Excel/CSV, párování podle CODE/EAN)"],
                               index={"direct": 0, "price_type": 1, "external": 2}.get(pcfg["mode"], 0), horizontal=True)
        prof["currency"] = st.text_input("Měna (CURRENCY)", value=prof.get("currency", "CZK"))
        if price_mode == "Přímé pole":
            sel = tag_select("Zdrojový tag ceny", pcfg.get("source", ""), "price_src")
            prof["price"] = {"mode": "direct", "source": "" if sel == "-- nevybráno --" else sel}
        elif price_mode.startswith("BMEcat"):
            attr = st.text_input("Atribut typu ceny (prázdné = ber první)", value=pcfg.get("price_type_attr", "price_type"))
            val = st.text_input("Hodnota atributu (např. net_list)", value=pcfg.get("price_type_value", "net_list"))
            sel = tag_select("Tag částky uvnitř ARTICLE_PRICE", pcfg.get("source", "PRICE_AMOUNT") or "PRICE_AMOUNT", "price_amt")
            prof["price"] = {"mode": "price_type", "source": "" if sel == "-- nevybráno --" else sel,
                              "price_type_attr": attr, "price_type_value": val}
        else:
            match_key = st.radio("Párovat podle", ["CODE", "EAN"], index=0 if pcfg.get("match_key", "CODE") == "CODE" else 1, horizontal=True)
            price_file = st.file_uploader("Ceník (Excel .xlsx nebo .csv)", type=["xlsx", "xlsm", "csv"], key="price_file_uploader")
            if price_file is not None:
                header, rows = read_table_file(price_file)
                if header:
                    c1, c2 = st.columns(2)
                    with c1:
                        id_col = st.selectbox(f"Sloupec s {match_key}", range(len(header)), format_func=lambda i: f"{header[i]} (sloupec {i+1})")
                    with c2:
                        price_col = st.selectbox("Sloupec s cenou", range(len(header)), index=min(1, len(header) - 1), format_func=lambda i: f"{header[i]} (sloupec {i+1})")
                    if st.button("Načíst ceník", type="primary"):
                        price_map = {}
                        for row in rows:
                            if id_col < len(row) and price_col < len(row) and row[id_col]:
                                price_map[normalize_id(row[id_col])] = row[price_col].replace(" ", "").replace(",", ".")
                        ss.external_prices = price_map
                        ss._ext_ver = ss.get("_ext_ver", 0) + 1
                        st.success(f"Načteno {len(price_map)} cen z ceníku.")
                        st.rerun()
            if ss.external_prices:
                st.caption(f"Aktuálně načteno **{len(ss.external_prices)}** cen v paměti.")
            prof["price"] = {"mode": "external", "match_key": match_key}

        st.markdown("#### Hmotnost a rozměry (nepovinné) → LOGISTIC")
        wcfg = prof["weight"]
        weight_mode = st.radio("Zdroj hmotnosti/rozměrů", ["Nevyplňovat", "Přímé pole (jen hmotnost)", "FEATURE podle názvu (BMEcat styl)"],
                                index={"none": 0, "direct": 1, "feature": 2}.get(wcfg["mode"], 0), horizontal=True)
        if weight_mode == "Nevyplňovat":
            prof["weight"] = {"mode": "none"}
        elif weight_mode == "Přímé pole (jen hmotnost)":
            sel = tag_select("Zdrojový tag hmotnosti", wcfg.get("source", ""), "weight_src")
            prof["weight"] = {"mode": "direct", "source": "" if sel == "-- nevybráno --" else sel}
        else:
            with st.expander("Nastavení FEATURE hmotnosti a rozměrů", expanded=True):
                st.caption(
                    "Shoptet používá pod LOGISTIC elementy WEIGHT / DEPTH / WIDTH / HEIGHT (LENGTH není "
                    "povolen). Rozměry appka přepočítá na cm (mm → cm), hmotnost z gramů na kg."
                )
                fe = st.text_input("Element vlastnosti", value=wcfg.get("feature_element", "FEATURE"))
                fn = st.text_input("Pole s názvem", value=wcfg.get("fname_field", "FNAME"))
                fv = st.text_input("Pole s hodnotou", value=wcfg.get("fvalue_field", "FVALUE"))
                fu = st.text_input("Pole s jednotkou", value=wcfg.get("funit_field", "FUNIT"))
                names = st.text_input("Název(y) = HMOTNOST (odděl čárkou)", value=wcfg.get("names", "Hmotnost, Weight"))
                c1, c2, c3 = st.columns(3)
                with c1:
                    width_names = st.text_input("Název(y) = ŠÍŘKA (nepovinné)", value=wcfg.get("width_names", ""))
                with c2:
                    height_names = st.text_input("Název(y) = VÝŠKA (nepovinné)", value=wcfg.get("height_names", ""))
                with c3:
                    length_names = st.text_input("Název(y) = HLOUBKA/DÉLKA → DEPTH (nepovinné)", value=wcfg.get("length_names", ""))
            prof["weight"] = {"mode": "feature", "feature_element": fe, "fname_field": fn, "fvalue_field": fv,
                               "funit_field": fu, "names": names, "width_names": width_names,
                               "height_names": height_names, "length_names": length_names}

    render_live_preview()
    nav_buttons()

# =============================================================================
# KROK 4 - KATEGORIE
# =============================================================================
elif step == 4:
    st.header(STEPS[3])
    prof = ss.profile
    ccfg = prof["category"]
    st.write("Necháváš produkty roztřídit do vlastního stromu kategorií ručně, takže nejpohodlnější "
             "je poslat je všechny do jedné (klidně skryté) kategorie typu 'Nezařazeno'.")
    mode = st.radio("Způsob nastavení kategorie", ["Pevná hodnota pro všechny produkty", "Vzít ze zdrojového pole"],
                     index=0 if ccfg["mode"] == "fixed" else 1)
    if mode.startswith("Pevná"):
        fixed = st.text_input("Kategorie pro všechny produkty", value=ccfg.get("fixed_value", "Nezařazeno"))
        prof["category"] = {"mode": "fixed", "fixed_value": fixed}
    else:
        root = ss.raw_root
        sample_items = root.xpath(f"//*[local-name()='{prof['item_tag']}']") if prof.get("item_tag") else []
        avail_tags = aggregate_descendant_tags(sample_items) if sample_items else []
        tag_options = ["-- nevybráno --"] + avail_tags
        cur = ccfg.get("source", "")
        idx = tag_options.index(cur) if cur in tag_options else 0
        sel = st.selectbox("Zdrojový tag kategorie", tag_options, index=idx)
        prof["category"] = {"mode": "direct", "source": "" if sel == "-- nevybráno --" else sel}

    render_live_preview()
    nav_buttons()

# =============================================================================
# KROK 5 - PARAMETRY
# =============================================================================
elif step == 5:
    st.header(STEPS[4])
    st.write(
        "Appka nejdřív navrhne, který tag v jedné 'vlastnosti' produktu (např. `<FEATURE>`) funguje jako "
        "rozlišovač druhu parametru (marketingový text / technický parametr / velikost), a najde jeho "
        "reálné hodnoty ve feedu. Ty pak každé hodnotě přiřadíš, co se s ní má stát - appka nic neuhodne napevno."
    )
    root = ss.raw_root
    prof = ss.profile
    item_tag = prof.get("item_tag", "")
    sample_items = root.xpath(f"//*[local-name()='{item_tag}']") if item_tag else []

    if not sample_items:
        st.warning("Nejdřív nastav element produktu v kroku 2 nebo 3.")
    else:
        pcfg = prof["params"]
        repeat_candidates = repeat_children_candidates(sample_items[0])
        repeat_names = [f"{t} ({c}×)" for t, c in repeat_candidates]
        repeat_map = {f"{t} ({c}×)": t for t, c in repeat_candidates}
        default_idx = next((i for i, (t, c) in enumerate(repeat_candidates) if t == pcfg.get("repeat_element")), 0)
        chosen_repeat = st.selectbox("Element jedné 'vlastnosti' produktu (např. FEATURE)",
                                      ["-- žádné parametry v tomto feedu --"] + repeat_names,
                                      index=default_idx + 1 if repeat_names else 0)
        if chosen_repeat == "-- žádné parametry v tomto feedu --":
            pcfg["repeat_element"] = ""
        else:
            pcfg["repeat_element"] = repeat_map.get(chosen_repeat, "")

        if pcfg["repeat_element"]:
            repeat_els = []
            for si in sample_items[:30]:
                repeat_els.extend(si.xpath(f".//*[local-name()='{pcfg['repeat_element']}']"))
            field_options = aggregate_descendant_tags(repeat_els) if repeat_els else []
            field_options = ["-- nevybráno --"] + sorted(set(field_options))

            def fsel(label, current, key):
                idx = field_options.index(current) if current in field_options else 0
                return st.selectbox(label, field_options, index=idx, key=key)

            c1, c2, c3 = st.columns(3)
            with c1:
                sel = fsel("Pole s NÁZVEM parametru (FNAME)", pcfg.get("fname_field", ""), "p_fname")
                pcfg["fname_field"] = "" if sel == "-- nevybráno --" else sel
            with c2:
                sel = fsel("Pole s HODNOTOU (FVALUE)", pcfg.get("fvalue_field", ""), "p_fvalue")
                pcfg["fvalue_field"] = "" if sel == "-- nevybráno --" else sel
            with c3:
                sel = fsel("Pole s JEDNOTKOU (nepovinné)", pcfg.get("funit_field", ""), "p_funit")
                pcfg["funit_field"] = "" if sel == "-- nevybráno --" else sel

            suggested = classifier_field_candidates(root, item_tag, pcfg["repeat_element"])
            scope_labels = {"parent": " (skupina vlastností - rodič)", "self": " (uvnitř jedné vlastnosti)"}
            classifier_options = [("none", "")] + suggested
            option_labels = ["-- neklasifikovat (nic negenerovat) --"] + \
                [f"{f}{scope_labels[s]}" for s, f in suggested]
            cur_scope, cur_field = pcfg.get("classifier_scope", ""), pcfg.get("classifier_field", "")
            cur_idx = next((i for i, (s, f) in enumerate(classifier_options) if s == cur_scope and f == cur_field), 0)
            chosen_idx = st.selectbox(
                "Rozlišovací pole (podle čeho appka pozná druh parametru) - navržené kandidáty jsou nahoře, "
                "'skupina vlastností' je pole nad celým blokem vlastností (typické pro Metabo apod.), "
                "'uvnitř vlastnosti' je pole přímo u jedné hodnoty (typické pro Bosch FDESCR)",
                range(len(option_labels)), index=cur_idx, format_func=lambda i: option_labels[i],
            )
            pcfg["classifier_scope"], pcfg["classifier_field"] = classifier_options[chosen_idx]

            if pcfg["classifier_field"]:
                values = get_classifier_values(root, item_tag, pcfg["repeat_element"], pcfg["classifier_scope"], pcfg["classifier_field"])
                st.markdown("#### Přiřaď nalezeným hodnotám, co se s nimi má stát")
                bucket_labels = {
                    "ignore": "Ignorovat",
                    "marketing": "Marketingový text → do popisu",
                    "packaging": "Obsah balení → úplný konec DESCRIPTION",
                    "technical": "Technický parametr → INFORMATION_PARAMETERS",
                    "size": "Velikost → PARAMETERS",
                }
                bucket_keys = list(bucket_labels.keys())
                for val, count in values:
                    default_bucket = pcfg["value_map"].get(val, guess_bucket(val))
                    sel = st.selectbox(
                        f"„{val}" + f"“  ({count}×)",
                        bucket_keys, index=bucket_keys.index(default_bucket),
                        format_func=lambda k: bucket_labels[k], key=f"bucket_{val}",
                    )
                    pcfg["value_map"][val] = sel

                if any(b == "marketing" for b in pcfg["value_map"].values()):
                    target = st.radio("Kam zapsat marketingové texty?",
                                       ["Do SHORT_DESCRIPTION", "Na konec DESCRIPTION (jako odrážky)"],
                                       index=0 if pcfg.get("marketing_target", "SHORT_DESCRIPTION") == "SHORT_DESCRIPTION" else 1)
                    pcfg["marketing_target"] = "SHORT_DESCRIPTION" if target.startswith("Do") else "DESCRIPTION"

    render_live_preview()
    nav_buttons()

# =============================================================================
# KROK 6 - OBRÁZKY
# =============================================================================
elif step == 6:
    st.header(STEPS[5])
    st.info(
        "**Doporučený postup:** nahraj obrázky do **Shoptet Správce souborů** (Administrace → Vzhled → "
        "Správce souborů). Shoptet si při importu obrázek stáhne, uloží na vlastní disk a připraví "
        "standardizované velikosti - obrázky pak žijou přímo v Shoptetu, ne u tebe. Jakmile jsou nahrané, "
        "zkopíruj URL té složky sem."
    )
    root = ss.raw_root
    prof = ss.profile
    item_tag = prof.get("item_tag", "")
    sample_items = root.xpath(f"//*[local-name()='{item_tag}']") if item_tag else []
    icfg = prof["images"]

    if sample_items:
        img_mode = st.radio("Zdroj obrázků ve vstupním feedu", ["Nevyplňovat", "Přímé pole (URL/název souboru)", "MIME_INFO/MIME (BMEcat styl)"],
                             index={"none": 0, "direct": 1, "mime": 2}.get(icfg["mode"], 0), horizontal=True)
        max_images = st.number_input("Max. počet obrázků na produkt (0 = bez omezení)", min_value=0,
                                      value=icfg.get("max_images", 8) or 0, step=1)
        add_desc = st.checkbox("Přidat atribut description (alt text) - Shoptet ho u obrázků používá",
                                value=icfg.get("add_description", True))
        avail_tags = aggregate_descendant_tags(sample_items)
        tag_options = ["-- nevybráno --"] + avail_tags

        def tag_select(label, current, key):
            idx = tag_options.index(current) if current in tag_options else 0
            return st.selectbox(label, tag_options, index=idx, key=key)

        preserved = {k: icfg.get(k) for k in ("public_base_url",) if icfg.get(k)}
        if img_mode == "Nevyplňovat":
            new_cfg = {"mode": "none", "max_images": max_images or None}
        elif img_mode.startswith("Přímé"):
            sel = tag_select("Zdrojový tag obrázku (může se opakovat)", icfg.get("source", ""), "img_src")
            new_cfg = {"mode": "direct", "source": "" if sel == "-- nevybráno --" else sel, "max_images": max_images or None}
        else:
            with st.expander("Nastavení MIME obrázků", expanded=True):
                me = st.text_input("Element jednoho obrázku", value=icfg.get("mime_element", "MIME"))
                pf = st.text_input("Pole účelu (prázdné = nefiltrovat)", value=icfg.get("purpose_field", "MIME_PURPOSE"))
                pv = st.text_input("Požadovaná hodnota účelu", value=icfg.get("purpose_value", "normal"))
                sf = st.text_input("Pole s názvem souboru", value=icfg.get("source_field", "MIME_SOURCE"))
            new_cfg = {"mode": "mime", "mime_element": me, "purpose_field": pf, "purpose_value": pv,
                       "source_field": sf, "max_images": max_images or None}
        new_cfg["add_description"] = add_desc
        new_cfg.update(preserved)
        prof["images"] = new_cfg
        icfg = prof["images"]

    st.markdown("---")
    shop_preview, _ = compute_shop()
    image_els = shop_preview.xpath(".//*[local-name()='IMAGE']") if shop_preview is not None else []
    raw_count = sum(1 for e in image_els if e.text and not e.text.strip().lower().startswith(("http://", "https://")))
    st.write(f"Obrázků celkem: **{len(image_els)}**, z toho bez platné URL: **{raw_count}**")
    st.caption(
        "Adresa cdn.myshoptet.com se vygeneruje sama až PO importu - do feedu nedávej vymyšlenou "
        "cdn.myshoptet.com URL, ale skutečnou veřejně dostupnou adresu, kam jsi obrázky přes FTP nahrál/-a."
    )
    base_url = st.text_input(
        "URL složky, kam jsi obrázky nahrál/-a přes FTP (appka připojí originální název souboru)",
        value=icfg.get("public_base_url") or "https://TVUJ-ESHOP.cz/...",
    )
    if st.button("Aplikovat URL na všechny obrázky bez platné adresy", type="primary"):
        prof["images"]["public_base_url"] = base_url
        st.success("URL složky uložena a použita na obrázky bez platné adresy.")
        st.rerun()

    st.markdown("#### 🔍 Test dostupnosti (ověří, jestli URL reálně vedou na obrázek)")
    if st.button("Zkontrolovat prvních pár obrázkových URL"):
        import requests
        shop_check, _ = compute_shop()
        check_urls = [e.text for e in shop_check.xpath(".//*[local-name()='IMAGE']")[:5] if e.text]
        for u in check_urls:
            try:
                r = requests.head(u, timeout=6, allow_redirects=True)
                ok = r.status_code == 200 and "image" in r.headers.get("Content-Type", "")
                (st.success if ok else st.error)(f"{r.status_code} — {u}")
            except Exception as e:
                st.error(f"Nedostupné ({e}) — {u}")

    with st.expander("🎥 Videa", expanded=False):
        st.caption(
            "Shoptet bere u videa jen KÓD YouTube videa: `<RELATED_VIDEO><YOUTUBE_VIDEO_CODE>kód</YOUTUBE_VIDEO_CODE>` "
            "(volitelně s `<TEXT>`). Appka kód vytáhne z odkazu (youtube.com/watch?v=..., youtu.be/..., embed, shorts). "
            "Odkazy mimo YouTube se přeskočí. Název obalového elementu (výchozí `RELATED_VIDEOS`) je můj odhad - "
            "validátor ti případně řekne, jak se jmenuje správně."
        )
        vcfg = prof.get("videos", {"mode": "none"})
        video_on = st.checkbox("Zahrnout videa z feedu (odkazy typu URL s popiskem 'video' nebo z YouTube)",
                                value=vcfg.get("mode") == "mime")
        if video_on:
            c1, c2 = st.columns(2)
            with c1:
                wrapper_tag = st.text_input("Obalovací element videí", value=vcfg.get("wrapper_tag") or "RELATED_VIDEOS")
            with c2:
                add_text = st.checkbox("Přidat <TEXT> (popisek videa = název produktu)", value=vcfg.get("add_text", False))
            prof["videos"] = {"mode": "mime", "mime_element": vcfg.get("mime_element", "MIME"),
                               "type_field": vcfg.get("type_field", "MIME_TYPE"), "type_value": vcfg.get("type_value", "url"),
                               "descr_field": vcfg.get("descr_field", "MIME_DESCR"), "descr_match": vcfg.get("descr_match", "video"),
                               "source_field": vcfg.get("source_field", "MIME_SOURCE"),
                               "wrapper_tag": wrapper_tag, "add_text": add_text}
        else:
            prof["videos"] = {"mode": "none"}

    with st.expander("🔗 Související a podobné produkty", expanded=False):
        st.caption(
            "⚠️ Přesný název tagů pro související produkty (výchozí `<RELATED_PRODUCTS><CODE>`) si "
            "ověř v Shoptet XML validátoru - appka je posílá takhle, ale nemám to z dokumentace "
            "stoprocentně potvrzené. Kódy odkazují na jiné produkty ve STEJNÉM feedu podle jejich CODE."
        )
        rcfg = prof.get("references", {"enabled": False, "included_types": []})
        item_tag_ref = prof.get("item_tag", "")
        sample_items_ref = ss.raw_root.xpath(f"//*[local-name()='{item_tag_ref}']") if item_tag_ref else []
        available_types = [t for t, c in get_attribute_values(ss.raw_root, item_tag_ref, "ARTICLE_REFERENCE", "type")] if sample_items_ref else []
        rcfg["enabled"] = st.checkbox("Zahrnout související/podobné produkty", value=rcfg.get("enabled", False))
        if rcfg["enabled"] and available_types:
            rcfg["included_types"] = st.multiselect(
                "Které typy reference zahrnout (nalezeno ve feedu)", available_types,
                default=rcfg.get("included_types") or available_types,
            )
            c1, c2 = st.columns(2)
            with c1:
                rcfg["wrapper_tag"] = st.text_input("Obalovací tag", value=rcfg.get("wrapper_tag", "RELATED_PRODUCTS"))
            with c2:
                rcfg["item_tag"] = st.text_input("Tag jednoho kódu", value=rcfg.get("item_tag", "CODE"))
        elif rcfg["enabled"]:
            st.info("Ve feedu jsem nenašel žádný ARTICLE_REFERENCE. Pokud se reference jmenují jinak, dej vědět.")
        prof["references"] = rcfg

    render_live_preview()
    nav_buttons()

# =============================================================================
# KROK 7 - KONTROLA A RUČNÍ OPRAVY
# =============================================================================
elif step == 7:
    st.header(STEPS[6])
    prof = ss.profile

    with st.expander("🏷️ Přidat text k názvu produktu (např. značku, pokud v NAME chybí)", expanded=False):
        st.caption("Aplikuje se na všechny produkty najednou, přímo v <NAME>.")
        c1, c2 = st.columns(2)
        with c1:
            prof["name_prefix"] = st.text_input("Text PŘED název (prefix)", value=prof.get("name_prefix", ""))
        with c2:
            prof["name_suffix"] = st.text_input("Text ZA název (suffix)", value=prof.get("name_suffix", ""))

    with st.expander("💰 Hromadná úprava ceny (NAŠE SLEVA + DPH)", expanded=False):
        prec_labels = [
            "Max. 2 desetinná místa (doporučeno - validátor Shoptetu povoluje nejvýš 2)",
            "Celá čísla (pozor: Shoptet k ceně sám přičte DPH, takže cena s DPH pak nemusí být celé číslo)",
        ]
        prec_choice = st.radio(
            "Přesnost <PRICE> a <STANDARD_PRICE> v exportu", prec_labels,
            index=0 if prof.get("price_precision", "2dp") == "2dp" else 1,
        )
        prof["price_precision"] = "2dp" if prec_choice == prec_labels[0] else "int"
        st.markdown("---")
        pa = prof.get("price_adjust", {"enabled": False, "source_mode": "net", "discount_percent": 0, "vat_percent": "21"})
        pa["enabled"] = st.checkbox("Zapnout přepočet ceny (cena s DPH na celé číslo + naše sleva)", value=pa.get("enabled", False))
        if pa["enabled"]:
            st.caption(
                "Cílem je celé číslo jako výsledná zobrazená cena s DPH v Shoptetu. Shoptet si k <PRICE> "
                "DPH dopočítává VŽDY sám (bere <PRICE> jako cenu bez DPH), takže appka do <PRICE> a "
                "<STANDARD_PRICE> zapíše cenu BEZ DPH - vypočítanou tak, aby po Shoptetově přičtení DPH "
                "vyšla co nejpřesněji kulatá částka. Pokud je zadaná i naše sleva, cena PŘED slevou jde do "
                "<STANDARD_PRICE> (přeškrtnutá cena v Shoptetu) a cena PO slevě do <PRICE>."
            )
            c1, c2 = st.columns(2)
            with c1:
                mode_choice = st.radio(
                    "Zdrojová cena z feedu je",
                    ["Bez DPH", "S DPH (např. z ceníku)"],
                    index=0 if pa.get("source_mode", "net") == "net" else 1,
                )
                pa["source_mode"] = "net" if mode_choice == "Bez DPH" else "gross"
                pa["vat_percent"] = st.text_input(
                    "Sazba DPH (%) - textově, ať jde snadno změnit při změně legislativy",
                    value=str(pa.get("vat_percent", "21")),
                )
            with c2:
                pa["discount_percent"] = st.number_input(
                    "NAŠE SLEVA (%) - odečte se od ceny s DPH", min_value=0.0, max_value=100.0,
                    value=float(pa.get("discount_percent", 0) or 0), step=0.5,
                )

        prof["price_adjust"] = pa

        if pa["enabled"]:
            shop_preview, _ = compute_shop()
            first_price = first_standard = None
            if shop_preview is not None:
                pe = shop_preview.xpath(".//*[local-name()='SHOPITEM'][1]/*[local-name()='PRICE']")
                if pe and pe[0].text:
                    first_price = pe[0].text
                std = shop_preview.xpath(".//*[local-name()='SHOPITEM'][1]/*[local-name()='STANDARD_PRICE']")
                first_standard = std[0].text if std and std[0].text else None
            if first_price:
                vat_frac = 1 + float(str(pa.get("vat_percent", "21")).replace(",", ".")) / 100.0
                displayed = float(first_price) * vat_frac
                line = f"Do feedu se zapíše `<PRICE>` = **{first_price}** (bez DPH) → v Shoptetu vyjde **{displayed:.2f} {prof.get('currency', '')}**"
                if first_standard:
                    displayed_std = float(first_standard) * vat_frac
                    line += f" (přeškrtnutá: ~~{displayed_std:.2f} {prof.get('currency', '')}~~)"
                st.caption(line)
                all_prices = [float(e.text) for e in shop_preview.xpath("./*[local-name()='SHOPITEM']/*[local-name()='PRICE']") if e.text]
                off = sum(1 for pv in all_prices if abs(round(pv * vat_frac, 2) - round(pv * vat_frac)) > 0.001)
                if all_prices:
                    st.caption(f"Cena s DPH vyjde přesně na celé číslo u **{len(all_prices) - off} z {len(all_prices)}** produktů "
                               f"(u zbylých {off} je to o haléř vedle - víc než 2 desetinná místa v <PRICE> validátor nepřijme).")

    with st.expander("📦 Viditelnost zásoby (dostupnost)", expanded=False):
        st.caption("Aplikuje se hromadně na všechny produkty.")
        av = prof.get("availability", {"enabled": False, "in_stock": "Skladem", "out_of_stock": "2-3 dny"})
        av["enabled"] = st.checkbox("Nastavit dostupnost hromadně", value=av.get("enabled", False))
        if av["enabled"]:
            c1, c2 = st.columns(2)
            with c1:
                av["in_stock"] = st.text_input(
                    "Text, když je SKLADEM (AVAILABILITY_IN_STOCK)", value=av.get("in_stock", "Skladem"),
                )
            with c2:
                idx = OUT_OF_STOCK_OPTIONS.index(av.get("out_of_stock")) if av.get("out_of_stock") in OUT_OF_STOCK_OPTIONS else 0
                av["out_of_stock"] = st.selectbox(
                    "Text, když NENÍ skladem (AVAILABILITY_OUT_OF_STOCK) - musí být jedna z předdefinovaných Shoptet hodnot",
                    OUT_OF_STOCK_OPTIONS, index=idx,
                )
        prof["availability"] = av

    shop, _ = compute_shop()
    issues = validate_shop(shop) if shop is not None else []
    n_mand = sum(1 for i in issues if i["mandatory"])
    n_opt = len(issues) - n_mand
    c1, c2, c3 = st.columns(3)
    c1.metric("Produktů", len(shop.xpath("./*[local-name()='SHOPITEM']")) if shop is not None else 0)
    c2.metric("🔴 Povinné chyby", n_mand)
    c3.metric("⚪ Doporučené", n_opt)

    if not issues:
        st.success("🎉 Feed neobsahuje žádné chybějící povinné ani doporučené elementy.")
    else:
        st.write("Nejdřív **povinné** chyby (musí se opravit, jinak Shoptet feed odmítne), pak doporučené; "
                  "uvnitř skupiny sestupně podle pozice v XML.")
        grouped = {}
        for i in issues:
            grouped.setdefault(i["tag"], []).append(i)
        group_order = sorted(grouped.keys(), key=lambda t: (not grouped[t][0]["mandatory"], -max(x["doc_pos"] for x in grouped[t])))
        for tag in group_order:
            group = grouped[tag]
            mand = group[0]["mandatory"]
            label = f"{'🔴 POVINNÉ' if mand else '⚪ Doporučené'} — <{tag}> chybí u {len(group)} produktů"
            with st.expander(label, expanded=mand):
                bulk_value = st.text_input(f"Hromadná hodnota pro všechny chybějící <{tag}>", key=f"bulk_{tag}")
                if st.button("Použít na všechny", key=f"apply_bulk_{tag}") and bulk_value:
                    for g in group:
                        ss.overrides.setdefault(g["code"], {})[tag] = bulk_value
                    st.success(f"Nastaveno u {len(group)} produktů.")
                    st.rerun()
                df = pd.DataFrame([{"CODE": g["code"], "NAME": g["name"], "Nová hodnota": ss.overrides.get(g["code"], {}).get(tag, "")} for g in group])
                edited = st.data_editor(df, key=f"editor_{tag}", use_container_width=True, disabled=["CODE", "NAME"])
                if st.button("Uložit individuální hodnoty výše", key=f"apply_ind_{tag}"):
                    cnt = 0
                    for _, row in edited.iterrows():
                        if row["Nová hodnota"]:
                            ss.overrides.setdefault(row["CODE"], {})[tag] = row["Nová hodnota"]
                            cnt += 1
                    if cnt:
                        st.success(f"Uloženo {cnt} hodnot.")
                        st.rerun()

    render_live_preview()
    nav_buttons()

# =============================================================================
# KROK 8 - EXPORT
# =============================================================================
elif step == 8:
    st.header(STEPS[7])
    shop, _ = compute_shop()
    export_shop = shop if shop is not None else etree.Element("SHOP")
    n_items = len(export_shop.xpath("./*[local-name()='SHOPITEM']"))
    st.caption(f"Do exportu jde {n_items} produktů (výběr z kroku 2 se uplatnil už na začátku).")

    issues = validate_shop(export_shop)
    n_mand = sum(1 for i in issues if i["mandatory"])
    if n_mand:
        st.error(f"⚠️ Pozor: export obsahuje {n_mand} povinných chyb - vrať se do kroku 7 a doplň je.")
    else:
        st.success("Export prochází kontrolou povinných elementů.")

    xml_bytes = shop_to_bytes(export_shop)
    st.download_button("📥 Stáhnout výsledný Shoptet XML feed", data=xml_bytes, file_name="shoptet_feed.xml",
                        mime="application/xml", type="primary")
    st.download_button(
        "💾 Stáhnout šablonu (.json) pro příští použití",
        data=json.dumps(ss.profile, ensure_ascii=False, indent=2).encode("utf-8"),
        file_name=f"{ss.profile.get('profile_name', 'sablona').replace(' ', '_')}.json",
        mime="application/json",
    )
    st.code(xml_bytes.decode("utf-8")[:20000], language="xml")
    if len(xml_bytes) > 20000:
        st.caption("(náhled zkrácen, stažený soubor je kompletní)")
    nav_buttons()
